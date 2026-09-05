import os
import uuid
import logging
from datetime import datetime
from pathlib import Path
from typing import Dict, Any, List, Optional

import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from src.config import settings
from src.agents.extractor import ExtractedInvoice
from src.agents.auditor import ComplianceReport

logger = logging.getLogger(__name__)


class LedgerManager:
    """Manages the official corporate expense audit ledger in formatted Excel."""

    def __init__(self, ledger_dir: Optional[Path] = None):
        self.ledger_dir = ledger_dir or settings.LEDGERS_DIR
        self.ledger_dir.mkdir(parents=True, exist_ok=True)
        self.ledger_path = self.ledger_dir / "expense_audit_ledger.xlsx"
        self._ensure_ledger_exists()

    def _ensure_ledger_exists(self):
        """Creates initialized workbook if not existing."""
        if not self.ledger_path.exists():
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Approved Expenses"

            headers = [
                "Transaction ID",
                "Date",
                "Submitter",
                "Vendor",
                "Category",
                "Currency",
                "Subtotal",
                "Tax",
                "Tip",
                "Total Amount",
                "Compliance Status",
                "Risk Score",
                "Decision",
                "Manager / Approver",
                "Decision Timestamp",
                "Notes / Flags",
            ]

            # Style Headers
            header_fill = PatternFill(start_color="1B365D", end_color="1B365D", fill_type="solid")
            header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
            border = Border(
                left=Side(style="thin", color="CCCCCC"),
                right=Side(style="thin", color="CCCCCC"),
                top=Side(style="thin", color="CCCCCC"),
                bottom=Side(style="thin", color="CCCCCC"),
            )

            for col_idx, h in enumerate(headers, start=1):
                cell = ws.cell(row=1, column=col_idx, value=h)
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.border = border

            ws.row_dimensions[1].height = 28
            wb.save(str(self.ledger_path))

    def record_expense(
        self,
        invoice: ExtractedInvoice,
        compliance: ComplianceReport,
        decision: str,
        approver: str = "Manager",
        submitter: str = "Employee",
        notes: str = "",
    ) -> Dict[str, Any]:
        """Appends a new audited expense entry to the ledger."""
        self._ensure_ledger_exists()
        tx_id = f"EXP-{datetime.now().strftime('%Y%m')}-{uuid.uuid4().hex[:5].upper()}"
        timestamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

        wb = openpyxl.load_workbook(str(self.ledger_path))
        ws = wb.active
        next_row = ws.max_row + 1

        # Determine decision tag
        decision_tag = decision.upper()
        flags_text = "; ".join(v.message for v in compliance.violations) if compliance.violations else "None"
        if notes:
            flags_text = f"{notes} | {flags_text}"

        row_values = [
            tx_id,
            invoice.date,
            submitter,
            invoice.vendor_name,
            invoice.category,
            invoice.currency,
            invoice.subtotal,
            invoice.tax,
            invoice.tip,
            invoice.total_amount,
            compliance.status,
            compliance.risk_score,
            decision_tag,
            approver,
            timestamp,
            flags_text,
        ]

        font_regular = Font(name="Segoe UI", size=10)
        border_cell = Border(
            left=Side(style="thin", color="E0E0E0"),
            right=Side(style="thin", color="E0E0E0"),
            top=Side(style="thin", color="E0E0E0"),
            bottom=Side(style="thin", color="E0E0E0"),
        )

        for col_idx, val in enumerate(row_values, start=1):
            cell = ws.cell(row=next_row, column=col_idx, value=val)
            cell.font = font_regular
            cell.border = border_cell
            cell.alignment = Alignment(vertical="center")

            # Number formatting
            if col_idx in (7, 8, 9, 10):  # Monies
                cell.number_format = "$#,##0.00"
                cell.alignment = Alignment(horizontal="right", vertical="center")
            elif col_idx in (2, 11, 12, 13, 15):
                cell.alignment = Alignment(horizontal="center", vertical="center")

        ws.row_dimensions[next_row].height = 22

        # Auto-adjust column widths
        for col in ws.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

        wb.save(str(self.ledger_path))
        download_url = f"{settings.PUBLIC_BASE_URL.rstrip('/')}/ledgers/{self.ledger_path.name}"

        return {
            "success": True,
            "transaction_id": tx_id,
            "filename": self.ledger_path.name,
            "filepath": str(self.ledger_path),
            "download_url": download_url,
            "row_index": next_row,
            "decision": decision_tag,
        }

    def get_summary_statistics(self) -> Dict[str, Any]:
        """Returns statistical totals from the expense ledger."""
        self._ensure_ledger_exists()
        try:
            df = pd.read_excel(str(self.ledger_path))
            if df.empty:
                return {"total_expenses": 0, "total_spend": 0.0, "approved_count": 0}

            total_spend = float(df[df["Decision"] == "APPROVED"]["Total Amount"].sum()) if "Total Amount" in df.columns else 0.0
            return {
                "total_records": len(df),
                "total_approved_spend_usd": round(total_spend, 2),
                "approved_count": int((df["Decision"] == "APPROVED").sum()) if "Decision" in df.columns else 0,
                "rejected_count": int((df["Decision"] == "REJECTED").sum()) if "Decision" in df.columns else 0,
            }
        except Exception as e:
            logger.error(f"Error reading summary stats: {e}")
            return {"error": str(e)}
